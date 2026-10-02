"""Exporta o quartil de produtividade para .xlsx, com fórmulas (muda o du_acc e tudo recalcula)."""
import io

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .services import QUARTIS

FONTE = "Arial"
CORES = {"Q1": "1A7F4B", "Q2": "5A9E2F", "Q3": "D08A10", "Q4": "C0392B"}
CLARAS = {"Q1": "D7EFE1", "Q2": "E3F0D6", "Q3": "F8E8C8", "Q4": "F5D5D1"}
CAB = PatternFill("solid", fgColor="1D2330")
BORDA = Border(bottom=Side(style="thin", color="E3E6EC"))


def _f(**kw):
    return Font(name=FONTE, **kw)


def _cabecalho(ws, linha, titulos, larguras):
    for col, (t, w) in enumerate(zip(titulos, larguras), start=1):
        c = ws.cell(row=linha, column=col, value=t)
        c.font = _f(bold=True, color="FFFFFF")
        c.fill = CAB
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = w


def gerar_xlsx(vendedores, sem_hierarquia, du_acc, dia_referencia, mes):
    wb = Workbook()

    # ------------------------------------------------------------ Vendedores
    ws = wb.active
    ws.title = "Vendedores"
    ws["A1"] = f"Quartil de produtividade FTTA — {mes:%m/%Y}"
    ws["A1"].font = _f(bold=True, size=14)
    ws["A2"] = "Produtividade = vendas ÷ du_acc · Q1 = 25% mais produtivos · empates ficam no mesmo quartil"
    ws["A2"].font = _f(italic=True, color="6B7385", size=9)

    ws["A3"] = "du_acc"
    ws["A3"].font = _f(bold=True)
    ws["B3"] = du_acc
    ws["B3"].font = _f(bold=True, color="0000FF")
    ws["B3"].fill = PatternFill("solid", fgColor="FFFF00")
    ws["B3"].comment = Comment(
        f"Fonte: dim_calendario (Databricks), du_acc em {dia_referencia:%d/%m/%Y}. "
        "Pode alterar: produtividade e quartis recalculam.", "painel")
    ws["C3"] = f"referência {dia_referencia:%d/%m/%Y}"
    ws["C3"].font = _f(color="6B7385", size=9)

    titulos = ["#", "Vendedor", "Gestor", "Vendas", "Vendas / DU", "Ranking", "Quartil"]
    _cabecalho(ws, 5, titulos, [6, 24, 22, 10, 13, 10, 10])
    ini = 6
    fim = ini + max(len(vendedores), 1) - 1
    faixa_prod = f"$E${ini}:$E${fim}"

    for i, l in enumerate(vendedores.itertuples(index=False), start=ini):
        ws.cell(row=i, column=1, value=i - ini + 1)
        ws.cell(row=i, column=2, value=l.username)
        ws.cell(row=i, column=3, value=l.gestor_direto)
        ws.cell(row=i, column=4, value=int(l.vendas))
        ws.cell(row=i, column=5, value=f"=IF($B$3>0,D{i}/$B$3,0)").number_format = "0.00"
        ws.cell(row=i, column=6, value=f'=COUNTIF({faixa_prod},">"&E{i})+1')
        ws.cell(row=i, column=7, value=f'="Q"&MIN(4,ROUNDUP(4*F{i}/COUNT({faixa_prod}),0))')
        for col in range(1, 8):
            c = ws.cell(row=i, column=col)
            c.font = _f()
            c.border = BORDA
            if col in (1, 4, 5, 6, 7):
                c.alignment = Alignment(horizontal="center" if col == 7 else "right")

    if len(vendedores):
        for q in QUARTIS:
            ws.conditional_formatting.add(
                f"G{ini}:G{fim}",
                FormulaRule(formula=[f'$G{ini}="{q}"'],
                            fill=PatternFill("solid", fgColor=CORES[q]),
                            font=Font(name=FONTE, bold=True, color="FFFFFF")))
    ws.freeze_panes = f"A{ini}"
    ws.auto_filter.ref = f"A5:G{fim}"

    # ------------------------------------------------------------ Resumo
    rs = wb.create_sheet("Resumo", 0)
    rs["A1"] = f"Resumo por quartil — {mes:%m/%Y}"
    rs["A1"].font = _f(bold=True, size=14)
    rs["A2"] = "du_acc"
    rs["A2"].font = _f(bold=True)
    rs["B2"] = "=Vendedores!B3"
    rs["B2"].font = _f(color="008000")
    _cabecalho(rs, 4, ["Quartil", "Vendedores", "Vendas", "Média vendas/DU", "Mín", "Máx"],
               [10, 12, 10, 17, 10, 10])
    g = f"Vendedores!$G${ini}:$G${fim}"
    d = f"Vendedores!$D${ini}:$D${fim}"
    e = f"Vendedores!$E${ini}:$E${fim}"
    for i, q in enumerate(QUARTIS, start=5):
        rs.cell(row=i, column=1, value=q)
        rs.cell(row=i, column=2, value=f'=COUNTIF({g},A{i})')
        rs.cell(row=i, column=3, value=f'=SUMIF({g},A{i},{d})')
        rs.cell(row=i, column=4, value=f'=IFERROR(AVERAGEIF({g},A{i},{e}),0)')
        rs.cell(row=i, column=5, value=f'=IF(B{i}=0,0,_xlfn.MINIFS({e},{g},A{i}))')
        rs.cell(row=i, column=6, value=f'=IF(B{i}=0,0,_xlfn.MAXIFS({e},{g},A{i}))')
        for col in range(1, 7):
            c = rs.cell(row=i, column=col)
            c.font = _f(bold=col == 1, color="FFFFFF" if col == 1 else "000000")
            c.fill = PatternFill("solid", fgColor=CORES[q] if col == 1 else CLARAS[q])
            c.alignment = Alignment(horizontal="center")
            if col >= 4:
                c.number_format = "0.00"
    rs["A9"] = "Total"
    rs["B9"] = "=SUM(B5:B8)"
    rs["C9"] = "=SUM(C5:C8)"
    rs["D9"] = f"=IFERROR(AVERAGE({e}),0)"
    rs["D9"].number_format = "0.00"
    for col in range(1, 5):
        rs.cell(row=9, column=col).font = _f(bold=True)
        rs.cell(row=9, column=col).alignment = Alignment(horizontal="center")

    # ------------------------------------------------------------ Fora da hierarquia
    fh = wb.create_sheet("Fora da hierarquia")
    fh["A1"] = "E-mails com venda no mês que não estão na tbhierarquia"
    fh["A1"].font = _f(bold=True)
    _cabecalho(fh, 3, ["E-mail vendedor", "Vendas"], [34, 10])
    for i, l in enumerate(sem_hierarquia.itertuples(index=False), start=4):
        fh.cell(row=i, column=1, value=l.email).font = _f()
        fh.cell(row=i, column=2, value=int(l.vendas)).font = _f()

    # openpyxl não grava valores calculados; o Excel recalcula tudo ao abrir.
    wb.calculation.fullCalcOnLoad = True

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
